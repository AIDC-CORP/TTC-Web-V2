import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def add_hyperlink(paragraph, url, text, color="0056B3", underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{r_id}"/>')
    new_run = parse_xml(f'<w:r {nsdecls("w", "r")}/>')
    rPr = parse_xml(f'<w:rPr {nsdecls("w", "r")}/>')
    if color:
        c = parse_xml(f'<w:color {nsdecls("w")} w:val="{color}"/>')
        rPr.append(c)
    if underline:
        u = parse_xml(f'<w:u {nsdecls("w")} w:val="single"/>')
        rPr.append(u)
    sz = parse_xml(f'<w:sz {nsdecls("w")} w:val="21"/>')
    rPr.append(sz)
    new_run.append(rPr)
    text_node = parse_xml(f'<w:t {nsdecls("w")} xml:space="preserve">{text}</w:t>')
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def parse_and_add_text(paragraph, text, font_size=10, color=RGBColor(0x33, 0x33, 0x33)):
    text = text.replace('<br>', '\n').replace('<br/>', '\n').strip()
    if not text:
        return
    
    # Check for website or email links to make them clickable
    link_match = re.search(r'(https?://[^\s]+|www\.[^\s]+|[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,})', text)
    if link_match and not text.startswith('|'):
        link_str = link_match.group(1)
        target_url = link_str if link_str.startswith('http') else (f'mailto:{link_str}' if '@' in link_str else f'http://{link_str}')
        start_idx, end_idx = link_match.span()
        
        prefix = text[:start_idx]
        if prefix:
            parse_and_add_text(paragraph, prefix, font_size, color)
            
        add_hyperlink(paragraph, target_url, link_str)
        
        suffix = text[end_idx:]
        if suffix:
            parse_and_add_text(paragraph, suffix, font_size, color)
        return

    pos = 0
    pattern = re.compile(r'\*\*(.*?)\*\*')
    for match in pattern.finditer(text):
        start, end = match.span()
        if start > pos:
            plain = text[pos:start].replace('*', '')
            if plain:
                r = paragraph.add_run(plain)
                r.font.size = Pt(font_size)
                r.font.color.rgb = color
        bold_txt = match.group(1).replace('*', '')
        if bold_txt:
            r = paragraph.add_run(bold_txt)
            r.bold = True
            r.font.size = Pt(font_size)
            r.font.color.rgb = color
        pos = end
    if pos < len(text):
        plain = text[pos:].replace('*', '')
        if plain:
            r = paragraph.add_run(plain)
            r.font.size = Pt(font_size)
            r.font.color.rgb = color

def convert_md_to_clean_docx(md_path, docx_path):
    doc = docx.Document()
    
    # Optimized margins (0.55 inch top/bottom, 0.6 inch left/right) for exact 2-page fit
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
        
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    normal_style.paragraph_format.space_after = Pt(1.5)
    normal_style.paragraph_format.line_spacing = 1.1

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_lines = []

    def process_table(t_lines):
        rows_data = []
        for tl in t_lines:
            tl_str = tl.strip()
            if tl_str.startswith('|') and tl_str.endswith('|'):
                parts = [p.strip() for p in tl_str.split('|')[1:-1]]
                if all(re.match(r'^:?-+:?$', p) for p in parts if p):
                    continue
                rows_data.append(parts)
        
        if not rows_data:
            return

        num_cols = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        for r_idx, r_data in enumerate(rows_data):
            row = table.rows[r_idx]
            
            # Prevent row from splitting across pages
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            
            if r_idx == 0:
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

            for c_idx, cell_text in enumerate(r_data):
                if c_idx < len(row.cells):
                    cell = row.cells[c_idx]
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.line_spacing = 1.05
                    if r_idx == 0:
                        set_cell_background(cell, "1B365D")
                        parse_and_add_text(p, cell_text, font_size=9.5, color=RGBColor(0xFF, 0xFF, 0xFF))
                    else:
                        bg_color = "F8F9FA" if r_idx % 2 == 1 else "FFFFFF"
                        set_cell_background(cell, bg_color)
                        parse_and_add_text(p, cell_text, font_size=9, color=RGBColor(0x22, 0x22, 0x22))

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(2)

    for line in lines:
        line_str = line.strip()

        if '<!-- pagebreak -->' in line_str or line_str == '\\pagebreak':
            doc.add_page_break()
            continue

        if line_str.startswith('|') and line_str.endswith('|'):
            in_table = True
            table_lines.append(line_str)
            continue
        else:
            if in_table:
                process_table(table_lines)
                in_table = False
                table_lines = []

        if not line_str:
            continue

        if line_str == '---':
            p_div = doc.add_paragraph()
            p_div.paragraph_format.space_before = Pt(2)
            p_div.paragraph_format.space_after = Pt(2)
            continue

        if line_str.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            parse_and_add_text(p, line_str[2:], font_size=15, color=RGBColor(0x1B, 0x36, 0x5D))
            p.runs[0].bold = True
        elif line_str.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(2)
            parse_and_add_text(p, line_str[3:], font_size=12, color=RGBColor(0x00, 0x56, 0xB3))
            p.runs[0].bold = True
        elif line_str.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1.5)
            parse_and_add_text(p, line_str[4:], font_size=10.5, color=RGBColor(0x1B, 0x36, 0x5D))
            p.runs[0].bold = True
        elif line_str.startswith('* ') or line_str.startswith('- '):
            item_text = line_str[2:].strip()
            if not item_text:
                continue
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.08
            parse_and_add_text(p, item_text, font_size=10)
        elif re.match(r'^\d+\.\s+', line_str):
            item_text = re.sub(r'^\d+\.\s+', '', line_str).strip()
            if not item_text:
                continue
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.08
            parse_and_add_text(p, item_text, font_size=10)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.08
            parse_and_add_text(p, line_str, font_size=10)

    if in_table:
        process_table(table_lines)

    try:
        doc.save(docx_path)
        print(f"Cleaned & Converted: {docx_path}")
    except PermissionError:
        alt_path = docx_path.replace(".docx", "-v2.docx")
        doc.save(alt_path)
        print(f"File locked by Word. Saved alternative to: {alt_path}")

if __name__ == "__main__":
    base_dir = r"d:\AIDC\TTC-V2\TTC-Web-V2\sales-kit"
    files = [
        r"01-capability-statement\01-capability-statement.md",
        r"02-corporate-brochure\02-corporate-brochure.md",
        r"03-executive-profile\03-executive-profile.md",
        r"04-presentation-pitchdeck\04-presentation-pitchdeck.md",
    ]
    for rel in files:
        md_p = os.path.join(base_dir, rel)
        docx_p = md_p.replace(".md", ".docx")
        if os.path.exists(md_p):
            convert_md_to_clean_docx(md_p, docx_p)
