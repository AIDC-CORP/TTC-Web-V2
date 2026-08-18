from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).with_name("01-capability-statement-final.docx")
LOGO = ROOT / "public" / "logo-ttc-removebg-DNXrVdJp.png"
HERO = ROOT / "public" / "about.png"
PROJECT_IMAGE = ROOT / "public" / "introsection1.jpg"

NAVY = "0B2D4D"
BLUE = "005BAC"
RED = "E31B23"
MID_GREY = "66717D"
TEXT = "24313D"
LINE = "D7DEE5"


def fontify(run, size=9.0, bold=False, color=TEXT, italic=False):
    run.font.name = "Arial"
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def format_paragraph(paragraph, before=0, after=0, line=1.04, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep
    return paragraph


def add_run(paragraph, text, size=9.0, bold=False, color=TEXT, italic=False):
    return fontify(paragraph.add_run(text), size, bold, color, italic)


def add_bottom_border(paragraph, color=LINE, size="8", space="4"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_hyperlink(paragraph, text, url, size=8.0, color=BLUE):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    run_properties.append(fonts)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    run_properties.append(color_el)
    size_el = OxmlElement("w:sz")
    size_el.set(qn("w:val"), str(int(size * 2)))
    run_properties.append(size_el)
    run.append(run_properties)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_header_footer(doc):
    """Keep repeating areas empty; the artifact renderer clips repeated content."""
    return


def add_page_branding(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    format_paragraph(p, after=0)
    add_run(p, "TTC", size=15.5, bold=True, color=RED)
    add_run(p, "  TÂN THÀNH CÔNG", size=8.2, bold=True, color=NAVY)
    add_bottom_border(p, color=BLUE, size="10", space="2")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    format_paragraph(p, before=0, after=3)
    add_run(p, "CAPABILITY STATEMENT 2026  |  VIETNAM", size=7.1, bold=True, color=MID_GREY)


def add_cropped_picture(paragraph, path, width_cm, height_cm, crop_top=0, crop_bottom=0, alt_text=""):
    shape = paragraph.add_run().add_picture(str(path), width=Cm(width_cm), height=Cm(height_cm))
    if alt_text:
        shape._inline.docPr.set("descr", alt_text)
    blip_fill = shape._inline.graphic.graphicData.pic.blipFill
    source_rect = OxmlElement("a:srcRect")
    source_rect.set("t", str(crop_top))
    source_rect.set("b", str(crop_bottom))
    blip_fill.insert(1, source_rect)
    return shape


def add_kicker(doc, text, before=4, after=1):
    p = doc.add_paragraph()
    format_paragraph(p, before=before, after=after, keep=True)
    add_run(p, text.upper(), size=7.2, bold=True, color=RED)
    return p


def add_heading(doc, text, before=0, after=3, size=13.0):
    p = doc.add_paragraph()
    format_paragraph(p, before=before, after=after, keep=True)
    add_run(p, text, size=size, bold=True, color=NAVY)
    add_bottom_border(p, color=LINE, size="5", space="2")
    return p


def add_capability(doc, number, title, body):
    p = doc.add_paragraph()
    format_paragraph(p, before=2.2, after=0.8, keep=True)
    add_run(p, number, size=8.2, bold=True, color=RED)
    add_run(p, "  " + title, size=10.0, bold=True, color=NAVY)
    p = doc.add_paragraph()
    format_paragraph(p, after=1.5, line=1.05)
    p.paragraph_format.left_indent = Mm(6)
    add_run(p, body, size=8.2, color=TEXT)


def add_bullet(doc, text, size=8.1, after=1.1):
    p = doc.add_paragraph()
    format_paragraph(p, after=after, line=1.03)
    p.paragraph_format.left_indent = Mm(5)
    p.paragraph_format.first_line_indent = Mm(-5)
    add_run(p, "■  ", size=6.0, bold=True, color=RED)
    add_run(p, text, size=size, color=TEXT)
    return p


def add_project(doc, name, sector, location, year, scale):
    p = doc.add_paragraph()
    format_paragraph(p, before=1.8, after=0.3, keep=True)
    add_run(p, name, size=9.0, bold=True, color=NAVY)
    p = doc.add_paragraph()
    format_paragraph(p, after=1.0, line=1.0)
    p.paragraph_format.left_indent = Mm(4)
    add_run(p, f"{sector}  |  {location}  |  {year}  |  Quy mô công bố: {scale}", size=7.6, color=MID_GREY)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(14)
    section.bottom_margin = Mm(14)
    section.left_margin = Mm(14)
    section.right_margin = Mm(14)
    section.header_distance = Mm(4)
    section.footer_distance = Mm(5)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    language = OxmlElement("w:lang")
    language.set(qn("w:val"), "vi-VN")
    language.set(qn("w:eastAsia"), "vi-VN")
    normal._element.rPr.append(language)
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.04

    add_header_footer(doc)
    add_page_branding(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, after=3)
    add_cropped_picture(
        p,
        HERO,
        18.2,
        5.8,
        crop_top=24000,
        crop_bottom=24000,
        alt_text="Minh họa phối hợp thiết kế và thi công công trình.",
    )

    add_kicker(doc, "Industrial construction  •  Design coordination  •  Project delivery", before=1, after=1)
    p = doc.add_paragraph()
    format_paragraph(p, after=2, line=0.98, keep=True)
    add_run(p, "TƯ VẤN • THIẾT KẾ • THI CÔNG", size=20.0, bold=True, color=NAVY)
    p = doc.add_paragraph()
    format_paragraph(p, after=3, keep=True)
    add_run(p, "CÔNG TRÌNH CÔNG NGHIỆP", size=20.0, bold=True, color=BLUE)

    p = doc.add_paragraph()
    format_paragraph(p, after=3, line=1.08)
    add_run(
        p,
        "Giải pháp tích hợp từ chuẩn bị đầu tư và thiết kế đa bộ môn đến thi công, cải tạo và hỗ trợ kỹ thuật sau bàn giao.",
        size=9.2,
        color=TEXT,
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, before=1, after=3)
    add_run(p, "2015 THÀNH LẬP", size=7.3, bold=True, color=BLUE)
    add_run(p, "   |   0107090447 MÃ SỐ DN   |   HÀ NỘI   |   04 NHÓM DỊCH VỤ", size=7.3, bold=True, color=MID_GREY)
    add_bottom_border(p, color=RED, size="7", space="4")

    p = doc.add_paragraph()
    format_paragraph(p, after=3, line=1.09)
    add_run(
        p,
        "Tân Thành Công cung cấp giải pháp cho dự án nhà máy và công trình công nghiệp tại Việt Nam. Chúng tôi tập trung vào tính khả thi khi thi công, sự phối hợp giữa các bộ môn và hiệu quả vận hành theo yêu cầu cụ thể của từng dự án.",
        size=8.8,
    )

    add_kicker(doc, "What we do", before=1)
    add_heading(doc, "Năng lực cốt lõi", after=1.5)
    add_capability(
        doc,
        "01",
        "Chuẩn bị đầu tư và phát triển dự án",
        "Khảo sát yêu cầu đầu tư; phối hợp quy hoạch, hồ sơ dự án, IRC, giấy phép xây dựng và môi trường theo phạm vi được giao; đánh giá phương án mặt bằng và khả năng triển khai.",
    )
    add_capability(
        doc,
        "02",
        "Thiết kế đa bộ môn và điều phối BIM",
        "Thiết kế kiến trúc, kết cấu và phối hợp MEP; rà soát xung đột và tính khả thi; ứng dụng 3D/4D/5D theo yêu cầu thông tin của dự án.",
    )
    add_capability(
        doc,
        "03",
        "Thi công công trình công nghiệp",
        "Nền móng, hạ tầng, kết cấu bê tông cốt thép, kết cấu thép tiền chế, mái và vách; điều phối MEP, QA/QC, HSE, nghiệm thu và bàn giao.",
    )
    add_capability(
        doc,
        "04",
        "Cải tạo và dịch vụ kỹ thuật",
        "Cải tạo, mở rộng nhà xưởng đang vận hành; nâng cấp công năng và hệ thống kỹ thuật; bảo hành, bảo trì và hỗ trợ sau bàn giao theo hợp đồng.",
    )

    add_kicker(doc, "How we deliver", before=2)
    add_heading(doc, "Phương pháp triển khai", after=1.5, size=11.6)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, after=2, line=1.0)
    add_run(p, "01 HIỂU YÊU CẦU  →  02 PHÁT TRIỂN GIẢI PHÁP  →  03 ĐIỀU PHỐI THIẾT KẾ  →  04 KIỂM SOÁT THI CÔNG  →  05 BÀN GIAO & HỖ TRỢ", size=6.9, bold=True, color=BLUE)

    add_kicker(doc, "Why TTC", before=2)
    add_heading(doc, "Giá trị khác biệt", after=1, size=11.6)
    add_bullet(doc, "Một đầu mối phối hợp giữa chủ đầu tư, tư vấn, nhà cung cấp và đội ngũ thi công.", size=7.8, after=0.5)
    add_bullet(doc, "Ưu tiên khả năng thi công để giảm rủi ro thay đổi tại công trường.", size=7.8, after=0.5)
    add_bullet(doc, "Giải pháp bám mục tiêu công năng, năng lượng, an toàn và khả năng mở rộng.", size=7.8, after=0)

    doc.add_page_break()

    add_page_branding(doc)
    add_kicker(doc, "Selected references", before=0)
    add_heading(doc, "Dự án tham chiếu", after=1, size=17.5)
    p = doc.add_paragraph()
    format_paragraph(p, after=2, line=1.06)
    add_run(p, "Một số dự án hiện được công bố trong danh mục năng lực của doanh nghiệp.", size=8.6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, after=2)
    add_cropped_picture(
        p,
        PROJECT_IMAGE,
        18.2,
        4.5,
        crop_top=29200,
        crop_bottom=29200,
        alt_text="Không gian bên trong một nhà xưởng công nghiệp đang hoàn thiện.",
    )

    add_project(doc, "Nhà máy Gạch Hồng Trang", "Vật liệu xây dựng", "Thái Nguyên", "2019", "100.000 m²")
    add_project(doc, "Nhà máy Japfa Com Feed Việt Nam", "Thức ăn chăn nuôi", "Thái Bình", "2018", "50.000 m²")
    add_project(doc, "Nhà máy A-One Timber Việt Nam", "Sản xuất gỗ", "Bắc Giang", "2018", "41.478 m²")
    add_project(doc, "Nhà máy Nội thất Fami", "Sản xuất nội thất", "Hưng Yên", "2016", "47.000 m²")
    add_project(doc, "Nhà máy Yusen", "Công nghiệp", "Hải Dương", "2017", "14.400 m²")

    p = doc.add_paragraph()
    format_paragraph(p, before=1, after=3, line=1.02)
    add_run(
        p,
        "Lưu ý: năm, địa điểm và quy mô theo danh mục dự án đang được công bố. Phạm vi hợp đồng và hồ sơ tham chiếu chi tiết được cung cấp trong hồ sơ dự thầu hoặc theo yêu cầu.",
        size=6.9,
        color=MID_GREY,
        italic=True,
    )

    add_kicker(doc, "Sector coverage", before=1)
    add_heading(doc, "Phân khúc phục vụ", after=1.5, size=11.6)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, after=3, line=1.0)
    add_run(p, "ĐIỆN TỬ & CƠ KHÍ  |  THỰC PHẨM & NÔNG NGHIỆP  |  LOGISTICS & KHO VẬN  |  VẬT LIỆU & NỘI THẤT  |  CẢI TẠO & MỞ RỘNG", size=6.9, bold=True, color=BLUE)

    add_kicker(doc, "Technical delivery", before=1)
    add_heading(doc, "Năng lực triển khai", after=1, size=11.6)
    p = doc.add_paragraph()
    format_paragraph(p, before=1, after=0.8, keep=True)
    add_run(p, "CÔNG TRÌNH & HỆ THỐNG", size=8.4, bold=True, color=NAVY)
    add_bullet(doc, "San nền, đường nội bộ, thoát nước, hạ tầng ngoài nhà; nền móng và bê tông cốt thép.", size=7.7, after=0.5)
    add_bullet(doc, "Kết cấu thép tiền chế, mái, vách, bao che công nghiệp và phối hợp MEP.", size=7.7, after=0.5)
    add_bullet(doc, "BIM coordination, kế hoạch thi công, nghiệm thu và hồ sơ bàn giao.", size=7.7, after=1.1)
    p = doc.add_paragraph()
    format_paragraph(p, after=0.8, keep=True)
    add_run(p, "QUẢN LÝ CHẤT LƯỢNG & AN TOÀN", size=8.4, bold=True, color=NAVY)
    add_bullet(doc, "Kế hoạch QA/QC, kiểm soát vật liệu, điểm kiểm soát và nghiệm thu.", size=7.7, after=0.5)
    add_bullet(doc, "Kế hoạch HSE theo yêu cầu dự án và quy định hiện hành.", size=7.7, after=0.5)
    add_bullet(doc, "Theo dõi tiến độ, thay đổi; hỗ trợ kỹ thuật, bảo hành và bảo trì theo hợp đồng.", size=7.7, after=1.5)

    add_kicker(doc, "Efficient & responsible solutions", before=1)
    add_heading(doc, "Hiệu quả và bền vững", after=1, size=11.6)
    p = doc.add_paragraph()
    format_paragraph(p, after=3, line=1.04)
    add_run(
        p,
        "Theo mục tiêu của chủ đầu tư: thông gió và chiếu sáng tự nhiên, vật liệu cách nhiệt, kiểm soát tải năng lượng, khả năng lắp đặt điện mặt trời và chuẩn bị hồ sơ phục vụ LEED, LOTUS hoặc EDGE khi thuộc phạm vi hợp đồng. Chứng nhận được xác lập cho từng dự án bởi tổ chức đánh giá độc lập.",
        size=7.7,
    )

    p = doc.add_paragraph()
    format_paragraph(p, before=2, after=1, keep=True)
    add_bottom_border(p, color=BLUE, size="12", space="3")
    add_run(p, "CÔNG TY CỔ PHẦN CÔNG NGHỆ XÂY DỰNG TÂN THÀNH CÔNG", size=9.0, bold=True, color=NAVY)
    p = doc.add_paragraph()
    format_paragraph(p, after=0.5)
    add_run(p, "TAN THANH CONG TECHNOLOGY CONSTRUCTION JOINT STOCK COMPANY", size=7.0, bold=True, color=MID_GREY)
    p = doc.add_paragraph()
    format_paragraph(p, after=0.5)
    add_run(p, "MST: 0107090447  |  Thành lập: 2015  |  +84 976 447 766", size=7.4, color=TEXT)
    p = doc.add_paragraph()
    format_paragraph(p, after=0.5)
    add_run(p, "Đăng ký: Số 39, ngõ 292, đường Kim Giang, Hà Nội  |  Văn phòng: Số 19N7B, KĐT Trung Hòa Nhân Chính, Hà Nội", size=6.9, color=MID_GREY)
    p = doc.add_paragraph()
    format_paragraph(p, after=0.5)
    add_hyperlink(p, "info@tanthanhcongjsc.com", "mailto:info@tanthanhcongjsc.com", size=7.3)
    add_run(p, "  |  ", size=7.3, color=MID_GREY)
    add_hyperlink(p, "tanthanhcongjsc.com", "https://tanthanhcongjsc.com", size=7.3)
    p = doc.add_paragraph()
    format_paragraph(p, before=0.5, after=0)
    add_run(p, "Hồ sơ pháp lý, chứng chỉ năng lực, tài liệu HSE và hồ sơ tham chiếu dự án được cung cấp theo yêu cầu và theo phạm vi áp dụng.", size=6.6, italic=True, color=MID_GREY)

    core = doc.core_properties
    core.title = "Capability Statement 2026 - Tân Thành Công JSC"
    core.subject = "Năng lực tư vấn, thiết kế và thi công công trình công nghiệp"
    core.author = "Tân Thành Công JSC"
    core.language = "vi-VN"
    core.keywords = "capability statement, industrial construction, BIM, Vietnam"
    core.comments = "Prepared for external business development use."

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
